"""
Script to generate the complete, professionally formatted Word Document (.docx)
for CHAPTER 11: APPLICATIONS for Smart Lab Management System (SLMS).
Strictly reflects the verified features and current operational state of the codebase.
"""

import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        run.font.name = "Calibri"
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        run.font.name = "Calibri"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        run.font.name = "Calibri"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x2E, 0x40, 0x53)
    return h

def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
    return p

def add_bullet_item(doc, bold_title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    r_bold = p.add_run(bold_title + ": ")
    r_bold.font.name = "Calibri"
    r_bold.font.size = Pt(11)
    r_bold.font.bold = True
    r_bold.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    return p

def build_chapter_11_docx(output_path: Path):
    doc = docx.Document()

    # Standard Academic Margins (1.25-inch left margin for binding, 1-inch top/bottom/right)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)

    # Document Header Title
    p_chap = doc.add_paragraph()
    p_chap.paragraph_format.space_before = Pt(0)
    p_chap.paragraph_format.space_after = Pt(2)
    p_chap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_chap_num = p_chap.add_run("CHAPTER 11\n")
    r_chap_num.font.name = "Calibri"
    r_chap_num.font.size = Pt(14)
    r_chap_num.font.bold = True
    r_chap_num.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    r_chap_title = p_chap.add_run("APPLICATIONS")
    r_chap_title.font.name = "Calibri"
    r_chap_title.font.size = Pt(20)
    r_chap_title.font.bold = True
    r_chap_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("―" * 45)
    r_div.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    # Intro Paragraph
    add_body_paragraph(
        doc,
        "This chapter outlines the practical, institutional, and operational applications of the Smart Lab Management System (SLMS). "
        "The system provides concrete solutions tailored to the operational realities of academic computer laboratories, polytechnic institutions, "
        "and university departments. Each application area described below directly leverages the verified capabilities implemented in the codebase, "
        "including real-time telemetry observation, dedicated port isolation, Global Remote Control batch execution, software inventory compliance, "
        "issue ticketing triage, hardware maintenance logging, and administrative audit trails."
    )

    # =========================================================================
    # 11.1 Academic Computer Laboratories & Practical Programming Classes
    # =========================================================================
    add_styled_heading(doc, "11.1 Academic Computer Laboratories & Practical Programming Classes", level=2)
    add_body_paragraph(
        doc,
        "The primary application domain of SLMS is in institutional computer laboratories (polytechnics, engineering colleges, and universities) "
        "where faculty members supervise dozens of students during practical laboratory coursework:"
    )
    add_bullet_item(doc, "Continuous Fleet Telemetry Observation", "Instructors can observe the operational status of all laboratory PCs from a centralized dashboard without pacing between aisles. Real-time gauges show which workstations are online, offline, or experiencing high CPU or memory load (RAM > 80% / 90%).")
    add_bullet_item(doc, "Non-Interfering Student Developer Environment", "Because SLMS host services bind exclusively to dedicated TCP port 18080, students learning web development, database programming, or API engineering (e.g., FastAPI, Django, Flask, or Node.js running on default ports 8000, 3000, or 5000) never encounter socket permission conflicts (WinError 10013).")
    add_bullet_item(doc, "Live Process Auditing", "Instructors can open the Computer Details page (/computers/:id) for any machine to inspect active processes (PIDs, CPU %, Memory share in MB). This read-only inspector enables faculty to verify that students are actively working on assigned coursework rather than running unapproved games or background tasks.")

    # =========================================================================
    # 11.2 Classroom Orchestration & Discipline via Global Remote Control
    # =========================================================================
    add_styled_heading(doc, "11.2 Classroom Orchestration & Discipline via Global Remote Control", level=2)
    add_body_paragraph(
        doc,
        "The Global Remote Control hub (/remote-control) empowers instructors to maintain classroom discipline and manage lab-wide power states "
        "through synchronized batch commands:"
    )
    add_bullet_item(doc, "Synchronized Broadcast Announcements", "Instructors can broadcast text announcements (up to 255 characters) simultaneously across selected student screens or the entire lab via SLMS_Notifier.exe (e.g., '10 minutes remaining to submit your lab exercise'). This eliminates the need to shout instructions across noisy computer halls.")
    add_bullet_item(doc, "Focusing Student Attention During Lectures", "During lecture demonstrations, whiteboard explanations, or projector presentations, the instructor can select all workstations and trigger 'Lock Workstation'. A fullscreen overlay freezes student monitors, preventing digital distractions while the instructor explains theoretical concepts.")
    add_bullet_item(doc, "Mass End-of-Day Power Management", "At the conclusion of practical classes, instructors can select all machines and execute 'Shutdown Computers' or 'Restart Computers' with a single click. This guarantees energy conservation and saves lab attendants from manually visiting and powering down 50+ computers individually.")

    # =========================================================================
    # 11.3 Software Inventory Auditing & Practical Exam Readiness
    # =========================================================================
    add_styled_heading(doc, "11.3 Software Inventory Auditing & Practical Exam Readiness", level=2)
    add_body_paragraph(
        doc,
        "The Software Inventory module (/software) solves the persistent administrative burden of verifying software prerequisites across lab computers:"
    )
    add_bullet_item(doc, "Pre-Exam Software Compliance Checks", "Ahead of university practical examinations (e.g., Python Lab, Java Lab, Database Lab, Web Technologies), faculty can search for specific software tools (e.g., Python 3.13, Visual Studio Code, PostgreSQL, Git) to confirm availability.")
    add_bullet_item(doc, "Workstation Installation Counts", "The system displays the exact count and hostnames of computers that have each application installed, instantly highlighting any workstation missing prerequisite compilers or libraries before students sit for an assessment.")
    add_bullet_item(doc, "Automated Registry Discovery", "Client agents automatically harvest software installation records from 32-bit and 64-bit Windows Registry uninstall trees, eliminating manual computer-by-computer software auditing.")

    # =========================================================================
    # 11.4 Laboratory Issue Tracking & Hardware Maintenance Management
    # =========================================================================
    add_styled_heading(doc, "11.4 Laboratory Issue Tracking & Hardware Maintenance Management", level=2)
    add_body_paragraph(
        doc,
        "SLMS replaces informal verbal reporting and paper logbooks with an integrated, digital laboratory maintenance system:"
    )
    add_bullet_item(doc, "Computer Defect Ticketing (/issues)", "Tracks hardware and software faults reported on specific workstations (e.g., faulty optical mouse, loose VGA cable, display flickering, network connectivity drops). Technicians can track issue progression across 'Open', 'In Progress', and 'Resolved' statuses.")
    add_bullet_item(doc, "Preventive Hardware Maintenance Records (/maintenance)", "Provides a structured digital log for technicians to document routine maintenance and hardware component servicing (e.g., RAM upgrades, CMOS battery replacements, thermal paste renewal, dust cleaning).")
    add_bullet_item(doc, "Equipment Lifecycle Management", "Builds a permanent, searchable hardware servicing history for each physical computer, supporting departmental warranty claims, Annual Maintenance Contract (AMC) reviews, and budget replacement planning.")

    # =========================================================================
    # 11.5 Administrative Audit Trails & Departmental Reporting
    # =========================================================================
    add_styled_heading(doc, "11.5 Administrative Audit Trails & Departmental Reporting", level=2)
    add_body_paragraph(
        doc,
        "The platform provides administrative accountability and compliance reporting for academic audits and departmental accreditation:"
    )
    add_bullet_item(doc, "Security Audit Logs (/audit)", "Maintains an immutable chronological record of every critical administrative operation performed on the platform, capturing the exact timestamp, operator username, action category (e.g., BATCH_COMMAND_LOCK, USER_LOGIN, SETTINGS_UPDATE), target workstation, and result status.")
    add_bullet_item(doc, "Institutional Data Export (/reports)", "Generates standardized CSV and JSON reports of laboratory computer inventories, installed software catalogs, and historical telemetry data for NBA/NAAC institutional accreditation, external inspections, and departmental reviews.")

    # Summary table
    add_styled_heading(doc, "Summary of Application Domains & Feature Mapping", level=3)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    hdr_cells[0].width = Inches(2.0)
    hdr_cells[1].width = Inches(2.2)
    hdr_cells[2].width = Inches(2.3)

    hdr_cells[0].text = "Application Domain"
    hdr_cells[1].text = "Primary SLMS Module Used"
    hdr_cells[2].text = "Operational Benefit Achieved"

    for cell in hdr_cells:
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        run = p.runs[0]
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    app_data = [
        (
            "Practical Lab Classes",
            "Dashboard & Computers (/computers)",
            "Real-time resource visibility without port collisions on developer ports."
        ),
        (
            "Classroom Orchestration",
            "Global Remote Control (/remote-control)",
            "Synchronized broadcast messaging, fullscreen screen locking, and mass shutdown."
        ),
        (
            "Exam Readiness Audits",
            "Software Inventory (/software)",
            "Instant verification of installed compilers and IDEs across all lab machines."
        ),
        (
            "Hardware Servicing",
            "Issues & Maintenance (/maintenance)",
            "Digital ticket triage and permanent servicing logs for AMC/warranty reviews."
        ),
        (
            "Departmental Accreditation",
            "Reports & Audit Logs (/reports, /audit)",
            "CSV/JSON export of hardware assets and chronological administrative action trails."
        )
    ]

    for item in app_data:
        row_cells = table.add_row().cells
        row_cells[0].width = Inches(2.0)
        row_cells[1].width = Inches(2.2)
        row_cells[2].width = Inches(2.3)

        for idx, text in enumerate(item):
            cell = row_cells[idx]
            set_cell_background(cell, "F9FAFB" if len(table.rows) % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
            if idx == 0:
                run.font.bold = True

    doc.save(str(output_path))
    print(f"[SUCCESS] Complete Chapter 11 Word document saved to: {output_path}")

if __name__ == "__main__":
    base_dir = Path(__file__).parent.parent
    docx_file = base_dir / "SLMS_Chapter_11_Applications.docx"
    build_chapter_11_docx(docx_file)
