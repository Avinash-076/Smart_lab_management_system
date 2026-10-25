"""
Script to generate the complete, professionally formatted Word Document (.docx)
for CHAPTER 8: OPERATION MANUAL for Smart Lab Management System (SLMS).
Strictly reflects the verified codebase features, navigation tabs, and remote commands.
"""

import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="2563EB"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        run.font.name = "Calibri"
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        run.font.name = "Calibri"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        run.font.name = "Calibri"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x2E, 0x40, 0x53)
    return h

def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
    return p

def add_bullet_item(doc, bold_title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    r_bold = p.add_run(bold_title + ": ")
    r_bold.font.name = "Calibri"
    r_bold.font.size = Pt(11)
    r_bold.font.bold = True
    r_bold.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    return p

def build_chapter_8_docx(output_path: Path):
    doc = docx.Document()

    # Standard Academic Margins (1.25-inch left margin for binding, 1-inch top/bottom/right)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)

    # Document Header Title
    p_chap = doc.add_paragraph()
    p_chap.paragraph_format.space_before = Pt(0)
    p_chap.paragraph_format.space_after = Pt(2)
    p_chap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_chap_num = p_chap.add_run("CHAPTER 8\n")
    r_chap_num.font.name = "Calibri"
    r_chap_num.font.size = Pt(14)
    r_chap_num.font.bold = True
    r_chap_num.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    r_chap_title = p_chap.add_run("OPERATION MANUAL")
    r_chap_title.font.name = "Calibri"
    r_chap_title.font.size = Pt(20)
    r_chap_title.font.bold = True
    r_chap_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("―" * 45)
    r_div.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    # Intro Paragraph
    add_body_paragraph(
        doc,
        "This chapter provides the operational manual for the Smart Lab Management System (SLMS). It details standard operational "
        "workflows for lab technicians, faculty instructors, and systems administrators. Topics covered include system startup, workstation "
        "enrollment, daily lab telemetry observation, Global Remote Control batch operations (message broadcasting, workstation locking, "
        "rebooting, and shutdown), software compliance auditing, issue management triage, maintenance logging, and administrative troubleshooting."
    )

    # =========================================================================
    # 8.1 System Startup and Workstation Onboarding
    # =========================================================================
    add_styled_heading(doc, "8.1 System Startup and Workstation Onboarding", level=2)
    add_body_paragraph(
        doc,
        "The SLMS Host service and client agents are engineered for minimal administrative overhead, automatically initializing core background "
        "services upon machine boot."
    )
    
    add_styled_heading(doc, "8.1.1 Launching the Host Application", level=3)
    add_bullet_item(doc, "Host Service Auto-Start", "The FastAPI backend operates under Windows Service Control Manager as SLMSHostService on dedicated port 18080. It starts automatically upon Host PC boot without requiring user logon.")
    add_bullet_item(doc, "Dashboard Authentication", "Launch the SLMS Host Dashboard from the desktop shortcut or Start Menu. Log in with your authorized administrator or instructor credentials. Upon validation, the dashboard establishes a persistent duplex WebSocket connection (ws://127.0.0.1:18080/ws/dashboard).")

    add_styled_heading(doc, "8.1.2 Workstation Enrollment Workflow", level=3)
    add_body_paragraph(
        doc,
        "Workstations must be enrolled once to establish cryptographically secured identity credentials with the Host:"
    )
    add_bullet_item(doc, "Step 1: Retrieve Enrollment Key", "Navigate to the Settings tab in the Host Dashboard. Note the Server IP address (e.g. 192.168.1.100) and generate an active 64-character Workstation Enrollment Key.")
    add_bullet_item(doc, "Step 2: Client Installation & Registration", "On the target student workstation, run the installation package or execute the headless CLI command with administrative privileges:")
    add_code_block(doc, 'SLMS_Client_Agent.exe enroll --url 192.168.1.100 --key "<ENROLLMENT_KEY>" --force\nSLMS_Client_Agent.exe install\nSLMS_Client_Agent.exe start')
    add_bullet_item(doc, "Step 3: Verification", "Port 18080 is automatically resolved. The agent encrypts its credentials via Windows DPAPI into %ProgramData%\\SLMS\\config\\service_credentials.enc and registers SLMSService in Session 0. The machine immediately appears in the live Computers list.")

    # =========================================================================
    # 8.2 Daily Laboratory Supervision & Hardware Monitoring
    # =========================================================================
    add_styled_heading(doc, "8.2 Daily Laboratory Supervision & Hardware Monitoring", level=2)
    
    add_styled_heading(doc, "8.2.1 Dashboard Fleet Overview (/dashboard)", level=3)
    add_body_paragraph(
        doc,
        "The Dashboard serves as the primary visual landing interface for instructors during practical classes, presenting four high-level metric cards:"
    )
    add_bullet_item(doc, "Total Computers", "Total number of registered machines in the laboratory asset inventory.")
    add_bullet_item(doc, "Online Workstations", "Workstations currently connected and transmitting telemetry within nominal parameters.")
    add_bullet_item(doc, "Offline Workstations", "Computers currently powered down, in sleep state, or disconnected from the local network.")
    add_bullet_item(doc, "Active Alerts", "Workstations exceeding defined operational thresholds (RAM > 90%, CPU > 95%, low disk headroom).")
    add_body_paragraph(
        doc,
        "Dynamic radial and linear progress meters summarize real-time average CPU load and physical memory consumption across the entire fleet."
    )

    add_styled_heading(doc, "8.2.2 Workstation Inspection & Process Monitoring (/computers)", level=3)
    add_body_paragraph(
        doc,
        "The Computers tab provides an interactive grid of workstation cards. Clicking any workstation card opens the Computer Details view (/computers/:id):"
    )
    add_bullet_item(doc, "Hardware Specifications", "Inspect operating system version, CPU architecture, installed RAM capacity, network IP/MAC addresses, and partition drive storage.")
    add_bullet_item(doc, "Live Process Inspector", "Displays a real-time list of all active processes executing on the student workstation, complete with Process ID (PID), executable name, CPU share percentage, and resident memory footprint (MB). This read-only inspector enables instructors to verify if required programming compilers or runaway scripts are active.")

    # =========================================================================
    # 8.3 Global Remote Control Operations
    # =========================================================================
    add_styled_heading(doc, "8.3 Global Remote Control Operations (/remote-control)", level=2)
    add_body_paragraph(
        doc,
        "The Global Remote Control hub empowers instructors to dispatch synchronized batch commands across multiple student workstations "
        "simultaneously over authenticated WebSockets."
    )

    add_styled_heading(doc, "8.3.1 Target Computer Selection", level=3)
    add_body_paragraph(
        doc,
        "In the Target Computer Table, instructors can filter machines by status, select individual workstations using dedicated checkboxes, "
        "or click 'Select All' to target the entire laboratory fleet. The selected summary counter updates instantly."
    )

    add_styled_heading(doc, "8.3.2 Command Center Action Suite", level=3)
    add_body_paragraph(
        doc,
        "The Command Center provides four distinct administrative actions configured with safety confirmation dialogs:"
    )
    add_bullet_item(doc, "1. Send Message (Action Card 1)", "Allows the instructor to type an announcement or instruction into the message text area (up to 255 characters). Clicking 'Send Message' dispatches the payload to target machines, displaying an instant notification banner on student screens via SLMS_Notifier.exe.")
    add_bullet_item(doc, "2. Lock Workstation (Action Card 2)", "Clicking 'Lock Workstations' prompts a confirmation modal. Once confirmed, a remote lock signal is transmitted, instantly freezing the selected student desktops with a fullscreen overlay to ensure undivided attention during classroom demonstrations.")
    add_bullet_item(doc, "3. Restart Computer (Action Card 3)", "Dispatches a clean operating system reboot signal across selected workstations, ideal for refreshing lab configurations between practical exam batches.")
    add_bullet_item(doc, "4. Shutdown Computer (Action Card 4)", "Allows instructors to power down the entire laboratory fleet cleanly at the end of the day, eliminating the need to physically visit each computer.")

    add_styled_heading(doc, "8.3.3 Execution Monitor & Command History", level=3)
    add_bullet_item(doc, "Execution Monitor", "Tracks real-time dispatch progression across target workstations, categorizing responses into pending, delivered, executed, or failed states.")
    add_bullet_item(doc, "Global Command History", "Maintains a permanent, searchable log of all previous commands. Includes date range filters (Today, Last 24 Hours, Last 7 Days, Last 30 Days) and a single-click 'Export to CSV' button for departmental audit records.")

    # =========================================================================
    # 8.4 Issue Management & Maintenance Records
    # =========================================================================
    add_styled_heading(doc, "8.4 Issue Management & Maintenance Records", level=2)
    
    add_styled_heading(doc, "8.4.1 Issue Management Workflow (/issues)", level=3)
    add_body_paragraph(
        doc,
        "The Issue Management module provides an integrated helpdesk for reporting and tracking computer hardware and software defects:"
    )
    add_bullet_item(doc, "Ticket Submission", "Captures workstation identifier, category (Hardware, Software, Network), severity level, student remarks, and timestamp.")
    add_bullet_item(doc, "Triage & Progression", "Technicians can review open tickets and update their operational status across 'Open', 'In Progress', and 'Resolved'.")
    add_bullet_item(doc, "Diagnostic Notes", "Allows technicians to record specific hardware faults (e.g., loose VGA cable, faulty optical mouse, corrupt compiler path).")

    add_styled_heading(doc, "8.4.2 Maintenance Records Module (/maintenance)", level=3)
    add_body_paragraph(
        doc,
        "The Maintenance module allows laboratory administrators to document routine preventive maintenance and hardware component replacements:"
    )
    add_bullet_item(doc, "Record Creation", "Log scheduled maintenance activities, specifying workstation ID, servicing date, and components serviced (e.g., thermal paste renewal, RAM upgrade, dust cleaning, OS re-imaging).")
    add_bullet_item(doc, "Audit Compliance", "Maintains an institutional maintenance history for each PC to support equipment lifecycle management and annual maintenance contract (AMC) reviews.")

    # =========================================================================
    # 8.5 Software Auditing, Reports & Security Audits
    # =========================================================================
    add_styled_heading(doc, "8.5 Software Auditing, Reports & Security Audits", level=2)

    add_styled_heading(doc, "8.5.1 Software Inventory Auditing (/software)", level=3)
    add_body_paragraph(
        doc,
        "Client agents automatically scan 32-bit and 64-bit Windows Registry paths to index all installed programs into a centralized catalog. "
        "Faculty can search for required development tools (e.g. Python, Visual Studio Code, PostgreSQL, Git) and verify the exact number "
        "and names of workstations where they are installed prior to practical exams."
    )

    add_styled_heading(doc, "8.5.2 Generating Reports (/reports)", level=3)
    add_body_paragraph(
        doc,
        "The Reports module provides comprehensive data export capabilities for academic and administrative audits:"
    )
    add_bullet_item(doc, "Inventory Reports", "Exports complete hardware and software profiles for all lab machines into structured CSV or JSON files.")
    add_bullet_item(doc, "Telemetry Logs", "Generates historical performance summaries detailing peak CPU/RAM utilization and uptime statistics.")

    add_styled_heading(doc, "8.5.3 Security Audit Logs (/audit)", level=3)
    add_body_paragraph(
        doc,
        "The Audit Logs view records every critical administrative operation performed on the platform, capturing the exact timestamp, "
        "operator username, action category (e.g. BATCH_COMMAND_LOCK, USER_LOGIN, SETTINGS_UPDATE), target workstation, and result status."
    )

    # =========================================================================
    # 8.6 Troubleshooting Common Operational Scenarios
    # =========================================================================
    add_styled_heading(doc, "8.6 Troubleshooting Common Operational Scenarios", level=2)
    add_body_paragraph(
        doc,
        "The following matrix summarizes common operational scenarios, diagnostic indicators, and verified resolution procedures:"
    )

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].width = Inches(1.8)
    hdr_cells[1].width = Inches(1.8)
    hdr_cells[2].width = Inches(2.9)

    hdr_cells[0].text = "Operational Scenario"
    hdr_cells[1].text = "Probable Root Cause"
    hdr_cells[2].text = "Verification & Resolution Procedure"

    for cell in hdr_cells:
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.runs[0]
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data Rows
    scenarios = [
        (
            "Workstation displays 'Offline' in Computers list",
            "Physical cable disconnected, Wi-Fi down, or SLMSService stopped.",
            "1. Verify physical Ethernet / Wi-Fi connection.\n2. In an admin terminal, run: sc query SLMSService\n3. If stopped, start via: SLMS_Client_Agent.exe start\n4. Inspect logs at: %ProgramData%\\SLMS\\logs\\agent.log"
        ),
        (
            "Batch Command shows 'pending' in Execution Monitor",
            "Target machine lost network connection before receiving WebSocket dispatch.",
            "1. Confirm target workstation is online.\n2. Verify Port 18080 inbound firewall rule is enabled on Host.\n3. The outbox/dispatcher will automatically retry upon machine reconnection."
        ),
        (
            "Client Enrollment Fails ('Server Unreachable')",
            "Host IP changed, Host Service stopped, or firewall blocking Port 18080.",
            "1. Verify Host IP using 'ipconfig' on Host PC.\n2. Confirm Host Service status: sc query SLMSHostService\n3. Add firewall rule: netsh advfirewall firewall add rule name=\"SLMS\" dir=in action=allow protocol=TCP localport=18080"
        ),
        (
            "Accidental Window Close on Host Dashboard",
            "Dashboard closed to System Tray (Close-to-Tray feature active).",
            "Double-click the SLMS icon in the Windows System Tray near the clock to restore the dashboard, or right-click the tray icon and select 'Quit SLMS Host' to fully terminate."
        ),
        (
            "Student Screen Lock Not Displaying",
            "SLMS_Notifier companion not running in Session 1 desktop session.",
            "1. Verify SLMS_Notifier.exe exists in HKLM Run registry key.\n2. Launch companion manually from: C:\\Program Files\\SLMS\\SLMS_Notifier.exe"
        )
    ]

    for item in scenarios:
        row_cells = table.add_row().cells
        row_cells[0].width = Inches(1.8)
        row_cells[1].width = Inches(1.8)
        row_cells[2].width = Inches(2.9)

        for idx, text in enumerate(item):
            cell = row_cells[idx]
            set_cell_background(cell, "F9FAFB" if len(table.rows) % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
            if idx == 0:
                run.font.bold = True

    # Trailing paragraph
    p_tail = doc.add_paragraph()
    p_tail.paragraph_format.space_before = Pt(12)
    p_tail.paragraph_format.space_after = Pt(6)
    r_tail = p_tail.add_run(
        "By adhering to these operational guidelines, laboratory staff can ensure maximum workstation uptime, enforce classroom "
        "discipline during practical sessions, maintain clean maintenance records, and resolve technical issues quickly and effectively."
    )
    r_tail.font.name = "Calibri"
    r_tail.font.size = Pt(11)
    r_tail.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

    doc.save(str(output_path))
    print(f"[SUCCESS] Complete Chapter 8 Word document saved to: {output_path}")

if __name__ == "__main__":
    base_dir = Path(__file__).parent.parent
    docx_file = base_dir / "SLMS_Chapter_8_Operation_Manual.docx"
    build_chapter_8_docx(docx_file)
