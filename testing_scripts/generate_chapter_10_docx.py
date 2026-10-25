"""
Script to generate the complete, professionally formatted Word Document (.docx)
for CHAPTER 10: SOURCE CODE & IMPLEMENTATION for Smart Lab Management System (SLMS).
Features designated file source boxes telling the user exactly which files to copy from.
"""

import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def add_code_file_box(doc, file_path, file_title, description, key_functions):
    """Add a professional framed callout box indicating the exact source code file."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC") # Soft Slate / Light Gray
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>' # Cyan/Blue accent border
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    
    r_hdr = p.add_run(f"📂 SOURCE CODE COMPONENT: {file_title}\n")
    r_hdr.font.name = "Calibri"
    r_hdr.font.size = Pt(11)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = RGBColor(0x03, 0x69, 0xA1) # Ocean Blue
    
    r_path_lbl = p.add_run("Exact File Path: ")
    r_path_lbl.font.name = "Calibri"
    r_path_lbl.font.size = Pt(10)
    r_path_lbl.font.bold = True
    
    r_path = p.add_run(f"{file_path}\n")
    r_path.font.name = "Consolas"
    r_path.font.size = Pt(9.5)
    r_path.font.bold = True
    r_path.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    r_desc_lbl = p.add_run("Component Description: ")
    r_desc_lbl.font.name = "Calibri"
    r_desc_lbl.font.size = Pt(10)
    r_desc_lbl.font.bold = True
    
    r_desc = p.add_run(f"{description}\n")
    r_desc.font.name = "Calibri"
    r_desc.font.size = Pt(10)
    r_desc.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    r_fn_lbl = p.add_run("Key Code Routines to Copy: ")
    r_fn_lbl.font.name = "Calibri"
    r_fn_lbl.font.size = Pt(10)
    r_fn_lbl.font.bold = True
    
    r_fn = p.add_run(f"{key_functions}\n\n")
    r_fn.font.name = "Consolas"
    r_fn.font.size = Pt(9)
    r_fn.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    r_note = p.add_run("[ PASTE YOUR CODE SNIPPET / LISTING BELOW THIS BOX ]")
    r_note.font.name = "Calibri"
    r_note.font.size = Pt(9.5)
    r_note.font.italic = True
    r_note.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Empty paragraph after table for spacing
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_image_placeholder(doc, figure_num, figure_title, description):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F9FAFB")
    set_cell_margins(cell, top=160, bottom=160, left=180, right=180)
    
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="9CA3AF"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="9CA3AF"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="9CA3AF"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="9CA3AF"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    
    r1 = p.add_run(f"📷 [ PASTE SCREENSHOT HERE FOR {figure_num} ]\n")
    r1.font.name = "Calibri"
    r1.font.size = Pt(10)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    
    r2 = p.add_run(f"({description})")
    r2.font.name = "Calibri"
    r2.font.size = Pt(9.5)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(10)
    
    r_cap_lbl = p_cap.add_run(f"{figure_num}: ")
    r_cap_lbl.font.name = "Calibri"
    r_cap_lbl.font.size = Pt(10)
    r_cap_lbl.font.bold = True
    r_cap_lbl.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    
    r_cap_txt = p_cap.add_run(figure_title)
    r_cap_txt.font.name = "Calibri"
    r_cap_txt.font.size = Pt(10)
    r_cap_txt.font.italic = True
    r_cap_txt.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="475569"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.0)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        run.font.name = "Calibri"
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        run.font.name = "Calibri"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        run.font.name = "Calibri"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x2E, 0x40, 0x53)
    return h

def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
    return p

def add_bullet_item(doc, bold_title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    r_bold = p.add_run(bold_title + ": ")
    r_bold.font.name = "Calibri"
    r_bold.font.size = Pt(11)
    r_bold.font.bold = True
    r_bold.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    return p

def build_chapter_10_docx(output_path: Path):
    doc = docx.Document()

    # Standard Academic Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)

    # Document Header Title
    p_chap = doc.add_paragraph()
    p_chap.paragraph_format.space_before = Pt(0)
    p_chap.paragraph_format.space_after = Pt(2)
    p_chap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_chap_num = p_chap.add_run("CHAPTER 10\n")
    r_chap_num.font.name = "Calibri"
    r_chap_num.font.size = Pt(14)
    r_chap_num.font.bold = True
    r_chap_num.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    r_chap_title = p_chap.add_run("SOURCE CODE & IMPLEMENTATION")
    r_chap_title.font.name = "Calibri"
    r_chap_title.font.size = Pt(20)
    r_chap_title.font.bold = True
    r_chap_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("―" * 45)
    r_div.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    # Intro Paragraph
    add_body_paragraph(
        doc,
        "This chapter details the source code architecture, component implementation, and file organization of the Smart Lab "
        "Management System (SLMS). Each subsection identifies the key source code files in the repository, their operational purpose, "
        "and designated source code reference boxes so code listings can be copied and reviewed directly from the implementation files."
    )

    # =========================================================================
    # 10.1 Project Directory Structure and Organization
    # =========================================================================
    add_styled_heading(doc, "10.1 Project Directory Structure and Organization", level=2)
    add_body_paragraph(
        doc,
        "The project is structured as a modular monorepo cleanly separating the Host FastAPI backend, the client agent background service, "
        "the Tauri desktop frontend, and the Inno Setup automated installers:"
    )

    dir_tree = (
        "Smart_lab_management_system/\n"
        "├── backend/                             # FastAPI Host Server (Port 18080)\n"
        "│   ├── alembic/                         # Database schema migrations\n"
        "│   ├── app/\n"
        "│   │   ├── models/                      # SQLAlchemy ORM models (Computer, Metrics, Commands)\n"
        "│   │   ├── routes/                      # API endpoint routers (Auth, Computers, Commands)\n"
        "│   │   ├── schemas/                     # Pydantic v2 validation models\n"
        "│   │   ├── services/                    # Business logic (Command, Alert, Report services)\n"
        "│   │   ├── websocket/                   # Real-time WebSocket connection managers\n"
        "│   │   ├── database.py                  # SQLite WAL engine configuration\n"
        "│   │   └── main.py                      # FastAPI application entrypoint\n"
        "│   ├── service_runner.py                # Host Windows SCM service runner\n"
        "│   └── pyproject.toml                   # Python dependencies managed via uv\n"
        "│\n"
        "├── client_agent/                        # Windows Client Agent & Service\n"
        "│   ├── core/                            # DPAPI security, outbox queue, runtime\n"
        "│   ├── modules/                         # Hardware, process, and software collectors\n"
        "│   ├── notifier/                        # Desktop notification companion (SLMS_Notifier)\n"
        "│   ├── server/                          # Enrollment and HTTP client\n"
        "│   ├── service/                         # SCM service implementation (SLMSService)\n"
        "│   └── main.py                          # Client CLI dispatcher\n"
        "│\n"
        "├── frontend/                            # React 19 + Tauri v2 Desktop Dashboard\n"
        "│   ├── src/                             # Pages, feature components, and hooks\n"
        "│   ├── src-tauri/                       # Native Rust desktop shell (tray, toasts)\n"
        "│   └── package.json                     # Frontend dependencies\n"
        "└── installer/                           # Inno Setup 6 packaging scripts"
    )
    add_code_block(doc, dir_tree)

    # =========================================================================
    # 10.2 Backend Implementation Using FastAPI
    # =========================================================================
    add_styled_heading(doc, "10.2 Backend Implementation Using FastAPI", level=2)
    add_body_paragraph(
        doc,
        "The Host Backend is implemented using FastAPI and served via Uvicorn. It exposes RESTful API endpoints for client telemetry "
        "ingestion, computer management, and full-duplex WebSockets for real-time dashboard events."
    )
    add_code_file_box(
        doc,
        file_path="backend/app/main.py",
        file_title="FastAPI Main Application Entry Point",
        description="Initializes the FastAPI application, registers global CORS middleware, defines the async lifespan event handler, and mounts all core API routers (auth, computers, commands, metrics, software, issues, reports).",
        key_functions="app = FastAPI(), lifespan(app: FastAPI), app.include_router(...)"
    )

    # =========================================================================
    # 10.3 Database Configuration, Models, and SQLite WAL Engine
    # =========================================================================
    add_styled_heading(doc, "10.3 Database Configuration, Models, and SQLite WAL Engine", level=2)
    add_body_paragraph(
        doc,
        "To guarantee high write concurrency from 100+ lab computers without database lockups, SLMS configures SQLite with Write-Ahead "
        "Logging (WAL mode), synchronous NORMAL, and a 30-second busy timeout."
    )
    add_code_file_box(
        doc,
        file_path="backend/app/database.py",
        file_title="Database Engine & SQLite WAL Concurrency PRAGMA",
        description="Configures SQLAlchemy engine, session maker, and event listener setting PRAGMA journal_mode=WAL, PRAGMA synchronous=NORMAL, and PRAGMA busy_timeout=30000.",
        key_functions="create_engine(), set_sqlite_pragma(), SessionLocal, Base"
    )
    add_code_file_box(
        doc,
        file_path="backend/app/models/remote_command.py",
        file_title="Remote Command ORM Models & Enums",
        description="Defines CommandType enum ('message', 'lock', 'restart', 'shutdown'), CommandStatus enum ('pending', 'delivered', 'executed', 'failed'), and RemoteCommand database model with foreign keys and timestamp indexes.",
        key_functions="class CommandType(str, Enum), class CommandStatus(str, Enum), class RemoteCommand(Base)"
    )

    # =========================================================================
    # 10.4 Client Agent Metric Probing and SQLite Outbox Implementation
    # =========================================================================
    add_styled_heading(doc, "10.4 Client Agent Metric Probing and SQLite Outbox Implementation", level=2)
    add_body_paragraph(
        doc,
        "The client agent runs non-blocking hardware probes every 10 seconds and buffers telemetry records atomically into an embedded "
        "local SQLite database to prevent data loss during network disruptions."
    )
    add_code_file_box(
        doc,
        file_path="client_agent/modules/hardware.py",
        file_title="Hardware Telemetry Probing Module",
        description="Uses psutil and Win32 system APIs to probe CPU usage percent, physical RAM usage/availability, and disk storage headroom without degrading workstation responsiveness.",
        key_functions="collect_system_metrics(), get_cpu_usage(), get_memory_info(), get_disk_partitions()"
    )
    add_code_file_box(
        doc,
        file_path="client_agent/core/outbox.py",
        file_title="Crash-Resilient Local SQLite Outbox Queue",
        description="Implements transactional outbox queue with atomic inserts, queue draining, and retry workers with exponential backoff for lossless offline buffering.",
        key_functions="class OutboxQueue: push(), pop_batch(), delete_dispatched(), get_queue_size()"
    )

    # =========================================================================
    # 10.5 Machine DPAPI Encryption & Automatic Port 18080 Normalization
    # =========================================================================
    add_styled_heading(doc, "10.5 Machine DPAPI Encryption & Automatic Port 18080 Normalization", level=2)
    add_body_paragraph(
        doc,
        "Client security routines prevent plaintext credential exposure on shared student PCs and normalize server addresses to dedicated port 18080."
    )
    add_code_file_box(
        doc,
        file_path="client_agent/core/security.py",
        file_title="DPAPI Machine Cryptography & URL Normalizer",
        description="Binds machine credentials using Windows DPAPI (CryptProtectData with CRYPTPROTECT_LOCAL_MACHINE flag) and automatically appends :18080 to raw IP inputs during enrollment.",
        key_functions="normalize_server_url(), encrypt_machine_secret(), decrypt_machine_secret()"
    )

    # =========================================================================
    # 10.6 Real-Time WebSocket Handlers and Remote Action Dispatchers
    # =========================================================================
    add_styled_heading(doc, "10.6 Real-Time WebSocket Handlers and Remote Action Dispatchers", level=2)
    add_body_paragraph(
        doc,
        "The command dispatcher orchestrates real-time command routing from the administration dashboard across persistent duplex WebSockets."
    )
    add_code_file_box(
        doc,
        file_path="backend/app/routes/command_router.py",
        file_title="Remote Command Batch Dispatch API",
        description="Handles batch command dispatching to selected workstations, captures execution duration, updates status transitions, and automatically writes audit log entries.",
        key_functions="@router.post('/batch'): issue_batch_commands(), get_fleet_overview(), get_command_history()"
    )
    add_code_file_box(
        doc,
        file_path="backend/app/websocket/websocket_manager.py",
        file_title="Duplex WebSocket Connection Manager",
        description="Maintains active client agent and dashboard WebSocket connections, routes bidirectional messages, and handles heartbeat keep-alive pings.",
        key_functions="connect_agent(), disconnect_agent(), broadcast_to_dashboard(), send_command_to_agent()"
    )

    # =========================================================================
    # 10.7 Frontend React 19 and Tauri v2 Rust Window Controls
    # =========================================================================
    add_styled_heading(doc, "10.7 Frontend React 19 and Tauri v2 Rust Window Controls", level=2)
    add_body_paragraph(
        doc,
        "The administration dashboard combines React 19 for modular component rendering with a Tauri v2 native Rust backend managing "
        "the window lifecycle, System Tray icon, and Windows Action Center toast notifications."
    )
    add_code_file_box(
        doc,
        file_path="frontend/src-tauri/src/lib.rs",
        file_title="Tauri v2 Native Rust Shell & Close-to-Tray Handler",
        description="Initializes Tauri plugins, configures the Windows System Tray menu ('Show Dashboard', 'Quit'), and intercepts WindowEvent::CloseRequested to hide the window instead of quitting.",
        key_functions="pub fn run(), on_window_event(|window, event| ... prevent_close() ... window.hide())"
    )
    add_code_file_box(
        doc,
        file_path="frontend/src/features/global-remote-control/components/CommandCenter/CommandCenter.jsx",
        file_title="Global Remote Control Command Center Component",
        description="Implements the 4 action cards (Send Message, Lock Workstations, Restart Computers, Shutdown Computers) with target summaries and safety confirmation dialogs.",
        key_functions="handleSendMessage(), handleConfirmAction(), onExecuteBatch(commandType, payload)"
    )

    # =========================================================================
    # 10.8 Automated Verification Test Suites (PyTest & Vitest)
    # =========================================================================
    add_styled_heading(doc, "10.8 Automated Verification Test Suites (PyTest & Vitest)", level=2)
    add_body_paragraph(
        doc,
        "Quality assurance is maintained through automated testing suites spanning backend unit tests, client security tests, and frontend integration tests:"
    )
    add_code_file_box(
        doc,
        file_path="client_agent/tests/test_security.py & backend/tests/",
        file_title="PyTest Security & Ingestion Verification Suite",
        description="Automated tests verifying URL normalization (port 18080 defaulting), DPAPI machine credential encryption/decryption, token validation, and SQLite WAL concurrent read/write behavior.",
        key_functions="test_normalize_server_url(), test_dpapi_machine_encryption(), test_batch_command_dispatch()"
    )
    add_code_file_box(
        doc,
        file_path="frontend/src/__tests__/",
        file_title="Vitest Frontend Component & Service Test Suite",
        description="Comprehensive suite of 271 passing tests across 22 test suites validating React 19 component rendering, authentication state transitions, and remote action modals.",
        key_functions="describe('CommandCenter'), describe('TargetComputerTable'), test('executes batch command')"
    )

    # =========================================================================
    # 10.9 Application Output Screens and Operational Proofs
    # =========================================================================
    add_styled_heading(doc, "10.9 Application Output Screens and Operational Proofs", level=2)
    add_body_paragraph(
        doc,
        "This section documents the visual output screens, graphical user interfaces, runtime terminal executions, "
        "and operational verification evidence of the Smart Lab Management System (SLMS). Each figure demonstrates "
        "a functioning component of the architecture, spanning client workstation background service execution, "
        "real-time fleet telemetry monitoring, software inventory compliance, global remote command dispatch, "
        "and active student session screen locking:"
    )

    add_styled_heading(doc, "Part 1: Packaging, Enrollment & System Service Verification", level=3)

    add_image_placeholder(
        doc,
        "Figure 10.1",
        "SLMS Host Inno Setup Wizard Installation Screen",
        "Screenshot showing SLMS_Host_Setup.exe completing setup on the teacher's PC"
    )

    add_image_placeholder(
        doc,
        "Figure 10.2",
        "Workstation Terminal Enrollment with Auto-Port 18080 Proof",
        "Screenshot of command prompt showing successful headless enrollment with DPAPI encryption"
    )

    add_image_placeholder(
        doc,
        "Figure 10.3",
        "Windows Service Control Manager (SCM) Execution Verification",
        "Screenshot of PowerShell running 'sc query SLMSService' showing STATE: 4 RUNNING"
    )

    add_styled_heading(doc, "Part 2: Working Graphical User Interface (GUI) Outputs", level=3)

    add_image_placeholder(
        doc,
        "Figure 10.4",
        "SLMS Central Laboratory Overview Dashboard and Telemetry Gauges",
        "Screenshot showing the main dashboard page (/dashboard) with summary statistics cards and average CPU/RAM gauges"
    )

    add_image_placeholder(
        doc,
        "Figure 10.5",
        "Live Workstation Matrix Grid displaying Real-Time Machine Status",
        "Screenshot showing the interactive grid of computer cards (/computers) with green, yellow, and red status badges"
    )

    add_image_placeholder(
        doc,
        "Figure 10.6",
        "Global Remote Control Command Center with Multi-Action Cards",
        "Screenshot showing the Remote Control hub (/remote-control) with Broadcast Message, Lock, Restart, and Shutdown cards"
    )

    add_image_placeholder(
        doc,
        "Figure 10.7",
        "Workstation Inspection Drawer with Live Process Table and Hardware Specs",
        "Screenshot showing the side drawer (/computers/:id) with hardware specs and active process list"
    )

    add_image_placeholder(
        doc,
        "Figure 10.8",
        "Centralized Software Inventory and Installed Application Audit View",
        "Screenshot showing the software inventory table (/software) with application names, versions, and workstation counts"
    )

    add_image_placeholder(
        doc,
        "Figure 10.9",
        "Laboratory Issue Ticketing and Problem Report Triage Interface",
        "Screenshot showing the issues table (/issues) with open/resolved tickets, machine IDs, and severity tags"
    )

    add_image_placeholder(
        doc,
        "Figure 10.10",
        "Student Workstation Full-Screen Lock Overlay Screen",
        "Screenshot of student PC locked by the teacher with the full-screen overlay rendered by SLMS_Notifier.exe"
    )

    add_image_placeholder(
        doc,
        "Figure 10.11",
        "Windows System Tray Integration and Native Desktop Notification Alert",
        "Screenshot showing taskbar system tray with SLMS icon and native Windows Action Center toast alert banner"
    )

    doc.save(str(output_path))
    print(f"[SUCCESS] Complete Chapter 10 Word document saved to: {output_path}")

if __name__ == "__main__":
    base_dir = Path(__file__).parent.parent
    docx_file = base_dir / "SLMS_Chapter_10_Source_Code_Guide.docx"
    try:
        build_chapter_10_docx(docx_file)
    except PermissionError:
        fallback_file = base_dir / "SLMS_Chapter_10_Source_Code_Guide_Updated.docx"
        print(f"[WARN] {docx_file} is locked by Word. Saving to {fallback_file}...")
        build_chapter_10_docx(fallback_file)
